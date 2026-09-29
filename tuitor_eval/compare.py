"""
Tuitor compare.py -- the regression decision for Tuitor, using exactly the main project's rules (evals/compare.py):
each metric's direction, gate/guardrail kind and tolerance come from evals/metric_registry.py.

    PASS    no gate blocked, no guardrail regressed. Safe to promote.
    REVIEW  a guardrail regressed beyond tolerance. A human decides.
    FAIL    a gate regressed (safety). Blocked.

Exit code: PASS=0, FAIL=1, REVIEW=2, so it can gate a script or CI.

    uv run python tuitor_eval/compare.py                               # baseline.json vs candidate.json in tuitor_eval/baselines/
    uv run python tuitor_eval/compare.py --baseline a.json --candidate b.json --all
"""

import argparse
import json
import shutil
import sys
from datetime import datetime

import common                                   # must be first: isolates this run inside tuitor_eval/
from evals.compare import EXIT_CODE, compare, load, print_report

BASELINE_NAME, CANDIDATE_NAME = "baseline.json", "candidate.json"


def baselines_dir():
    return common.HERE / "baselines"


def save_snapshot(snapshot, name):
    """Write baselines/<name>.json, plus a timestamped copy in baselines/history/, so overwriting never loses an earlier run."""
    folder = baselines_dir()
    (folder / "history").mkdir(parents=True, exist_ok=True)
    path = folder / f"{name}.json"
    path.write_text(json.dumps(snapshot, indent=2, sort_keys=True), encoding="utf-8")
    shutil.copy2(path, folder / "history" / f"{name}_{datetime.now().strftime('%Y-%m-%d_%H%M%S')}.json")
    return str(path)


def describe(snapshot, fallback):
    meta = snapshot.get("metadata", {})
    return f"{meta.get('label') or fallback}   (prompt {meta.get('prompt_hash', '?')}, judge {meta.get('judge_model', '?')}, {meta.get('created_at', '?')})"


def add_regression_sheet(xlsx_path, verdict, rows):
    """Append a 'Regression check' sheet to an Excel report: verdict, then every metric with baseline, candidate, delta and status."""
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = load_workbook(xlsx_path)
    ws = wb.create_sheet("Regression check", 1)
    head_fill, head_font = PatternFill("solid", fgColor="1F3864"), Font(name="Arial", bold=True, color="FFFFFF", size=10)
    ws["A1"], ws["B1"] = "Verdict", verdict
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    ws["B1"].font = Font(name="Arial", bold=True, size=12,
                         color={"PASS": "375623", "REVIEW": "7F6000", "FAIL": "9C0006"}[verdict])
    ws["A2"] = ("Compared with the baseline using the rules in evals/metric_registry.py (gate = must not regress, guardrail = "
                "regression is flagged for review, info = shown only). Delta is candidate minus baseline.")
    ws["A2"].font = Font(name="Arial", italic=True, size=9, color="595959")
    for i, h in enumerate(["Metric id", "Kind", "Better", "Baseline", "Candidate", "Delta", "Status"], start=1):
        c = ws.cell(4, i, h)
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center")
    order = {"blocked": 0, "regressed": 1, "dropped": 2, "new": 3, "improved": 4, "flat": 5, "info": 6}
    for r, row in enumerate(sorted(rows, key=lambda x: (order.get(x["status"], 9), x["id"])), start=5):
        to_cell = lambda v: (int(v) if isinstance(v, bool) else v)
        for c, v in enumerate([row["id"], row["kind"], row["direction"], to_cell(row["baseline"]), to_cell(row["candidate"])], start=1):
            ws.cell(r, c, v).font = Font(name="Arial", size=10)
        ws.cell(r, 6, f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},"")')
        ws.cell(r, 7, row["status"]).font = Font(name="Arial", size=10, bold=True)
        for c in (4, 5, 6):
            ws.cell(r, c).number_format = "0.000"
    last = 4 + len(rows)
    bad, good = PatternFill("solid", bgColor="F8CBAD"), PatternFill("solid", bgColor="C6E0B4")
    for word in ("blocked", "regressed"):
        ws.conditional_formatting.add(f"G5:G{last}", CellIsRule(operator="equal", formula=[f'"{word}"'], fill=bad, font=Font(name="Arial", bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"G5:G{last}", CellIsRule(operator="equal", formula=['"improved"'], fill=good, font=Font(name="Arial", bold=True, color="375623")))
    for col, w in zip("ABCDEFG", [46, 11, 8, 12, 12, 11, 11]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wb.save(xlsx_path)


def main():
    parser = argparse.ArgumentParser(description="Compare two Tuitor snapshots and return a verdict (same rules as evals/compare.py).")
    parser.add_argument("--baseline", default=None, help=f"default: tuitor_eval/baselines/{BASELINE_NAME}")
    parser.add_argument("--candidate", default=None, help=f"default: tuitor_eval/baselines/{CANDIDATE_NAME}")
    parser.add_argument("--all", action="store_true", help="also show info metrics")
    args = parser.parse_args()

    baseline_path = args.baseline or str(baselines_dir() / BASELINE_NAME)
    candidate_path = args.candidate or str(baselines_dir() / CANDIDATE_NAME)
    baseline, candidate = load(baseline_path), load(candidate_path)
    print(f"baseline  : {describe(baseline, baseline_path)}")
    print(f"candidate : {describe(candidate, candidate_path)}")
    verdict, rows = compare(baseline, candidate)
    print_report(verdict, rows, show_all=args.all)
    sys.exit(EXIT_CODE[verdict])


if __name__ == "__main__":
    main()
