"""
report_xlsx.py -- the full Excel report written by `python -m evals.run_suite`.

Every suite run writes ONE NEW workbook, reports/suite_<local date>_<time>.xlsx, and never overwrites an
earlier one. (The cumulative tracker in reports/eval_tracker.xlsx is separate: see evals/tracker.py.)

Sheets
  Run info    what was measured and with which settings (top_k, models, chunking, store size, git/prompt ids)
  Metrics     EVERY metric of the run, including the 'info' ones that the trimmed snapshot JSON leaves out
  By metric   per-metric case counts / pass rate / average score -- FORMULAS over the Cases sheet
  Ops report  latency / cost / reliability tables, pulled from Metrics by formula, with PASS/FAIL against the targets
  Cases       one row per (case, metric): the question, the app's answer, the expected answer, score, verdict and
              the judge's reason -- for every DeepEval-based eval (retriever, generator, pipeline, application, safety)

Per-case rows are captured by wrapping each eval module's `evaluate` for the duration of the run
(CaseCollector.capturing), so the eval files themselves do not change.
"""

import re
from contextlib import contextmanager

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from evals.metric_registry import rule_for
from evals.tracker import FONT, _metric_sort_key, _nice_name, _number_format, _style

MAX_CELL = 30000                                        # Excel's hard limit is 32,767 characters per cell
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")   # control characters openpyxl refuses to write


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        value = "\n---\n".join(str(v) for v in value)
    s = _ILLEGAL.sub("", str(value))
    return s if len(s) <= MAX_CELL else s[:MAX_CELL] + " ...[truncated]"


# ------------------------------------------------------------------ capturing per-case results
def extract_cases(group, result):
    """A DeepEval EvaluationResult -> one row per (case, metric)."""
    tests = getattr(result, "test_results", None)
    if tests is None:
        tests = result if isinstance(result, list) else []
    rows = []
    for i, tr in enumerate(tests, start=1):
        for m in (getattr(tr, "metrics_data", None) or getattr(tr, "metrics", None) or []):
            error = getattr(m, "error", None)
            rows.append({
                "group": group,
                "case": i,
                "metric": getattr(m, "name", "unknown"),
                "result": "ERROR" if error else ("PASS" if getattr(m, "success", False) else "FAIL"),
                "score": getattr(m, "score", None),
                "threshold": getattr(m, "threshold", None),
                "input": _text(getattr(tr, "input", "")),
                "actual": _text(getattr(tr, "actual_output", "")),
                "expected": _text(getattr(tr, "expected_output", "")),
                "reason": _text(error or getattr(m, "reason", "")),
                "context": _text(getattr(tr, "retrieval_context", None)),
                "judge": _text(getattr(m, "evaluation_model", "")),
            })
    return rows


class CaseCollector:
    def __init__(self):
        self.rows = []

    @contextmanager
    def capturing(self, targets):
        """targets = [(group, module)]. Temporarily wraps module.evaluate so each DeepEval result is recorded."""
        patched = []
        try:
            for group, module in targets:
                original = module.evaluate
                patched.append((module, original))
                module.evaluate = self._wrap(group, original)
            yield self
        finally:
            for module, original in patched:
                module.evaluate = original

    def _wrap(self, group, evaluate_fn):
        def wrapper(*args, **kwargs):
            result = evaluate_fn(*args, **kwargs)
            try:
                self.rows.extend(extract_cases(group, result))
            except Exception as e:      # the eval itself succeeded; never lose it over the report
                print(f"[report] could not read per-case results for '{group}': {type(e).__name__}: {e}")
            return result
        return wrapper


# ------------------------------------------------------------------ the workbook
def write_suite_report(path, metrics, cases, info, targets):
    """metrics {id: number}; cases [extract_cases rows]; info [(key, value)]; targets {name: number}."""
    from pathlib import Path

    wb = Workbook()
    wb.remove(wb.active)
    _sheet_info(wb, info)
    _sheet_metrics(wb, metrics)
    _sheet_by_metric(wb, cases)
    _sheet_ops(wb, len(metrics), targets)
    _sheet_cases(wb, cases)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _header(ws, row, headers):
    st = _style()
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row, i, h)
        c.font, c.fill = st["head_font"], st["head_fill"]
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _pass_fail_formatting(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=PatternFill("solid", bgColor="C6E0B4"), font=Font(name=FONT, color="375623")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=PatternFill("solid", bgColor="F8CBAD"), font=Font(name=FONT, bold=True, color="9C0006")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ERROR"'], fill=PatternFill("solid", bgColor="FFE699"), font=Font(name=FONT, bold=True, color="7F6000")))


def _sheet_info(wb, info):
    st = _style()
    ws = wb.create_sheet("Run info")
    ws["A1"] = "Eval suite report"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    for r, (key, value) in enumerate(info, start=3):
        ws.cell(r, 1, key).font = st["bold"]
        c = ws.cell(r, 2, value)
        c.font, c.alignment = st["body"], Alignment(wrap_text=True, vertical="top")
    end = 3 + len(info) + 1
    ws.cell(end, 1, "Sheets").font = st["bold"]
    ws.cell(end, 2, "Metrics = every number of the run.  By metric = pass counts computed from Cases.  Ops report = latency / cost / "
                    "reliability tables with PASS/FAIL.  Cases = each question, answer, score and the judge's reason.").font = st["body"]
    ws.cell(end, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100


def _sheet_metrics(wb, metrics):
    st = _style()
    ws = wb.create_sheet("Metrics")
    _header(ws, 1, ["Group", "Metric id", "Metric", "Stat", "Value", "Kind", "Better", "Tol", "Rel tol"])
    for r, mid in enumerate(sorted(metrics, key=_metric_sort_key), start=2):
        group, name, stat = _nice_name(mid)
        rule = rule_for(mid)
        for c, v in enumerate([group, mid, name, stat, metrics[mid], rule["kind"], rule["direction"], rule["tol"], rule["rel_tol"]], start=1):
            ws.cell(r, c, v).font = st["body"]
        ws.cell(r, 5).number_format = _number_format(mid)
        ws.cell(r, 5).font = st["bold"]
    for col, w in zip("ABCDEFGHI", [12, 46, 26, 22, 14, 10, 8, 6, 8]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(len(metrics) + 1, 2)}"


def _sheet_cases(wb, cases):
    st = _style()
    ws = wb.create_sheet("Cases")
    _header(ws, 1, ["Group", "Case", "Metric", "Result", "Score", "Threshold", "Input", "Actual output", "Expected output",
                    "Judge's reason", "Retrieved context", "Judge model"])
    for r, row in enumerate(cases, start=2):
        values = [row["group"], row["case"], row["metric"], row["result"], row["score"], row["threshold"], row["input"], row["actual"],
                  row["expected"], row["reason"], row["context"], row["judge"]]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.font = st["body"]
            cell.alignment = Alignment(vertical="top", wrap_text=c >= 7)
        ws.cell(r, 5).number_format = ws.cell(r, 6).number_format = "0.000"
    for col, w in zip("ABCDEFGHIJKL", [12, 6, 24, 9, 8, 9, 40, 60, 40, 60, 60, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:L{max(len(cases) + 1, 2)}"
    _pass_fail_formatting(ws, f"D2:D{max(len(cases) + 1, 2)}")


def _sheet_by_metric(wb, cases):
    st = _style()
    ws = wb.create_sheet("By metric", 2)
    _header(ws, 1, ["Group", "Metric", "Cases", "Passed", "Failed", "Pass rate", "Avg score"])
    last = max(len(cases) + 1, 2)
    pairs = sorted({(c["group"], c["metric"]) for c in cases}, key=lambda p: (_metric_sort_key(p[0] + ".x")[0], p[0], p[1]))
    for r, (group, metric) in enumerate(pairs, start=2):
        crit = f"Cases!$A$2:$A${last},$A{r},Cases!$C$2:$C${last},$B{r}"
        ws.cell(r, 1, group)
        ws.cell(r, 2, metric)
        ws.cell(r, 3, f"=COUNTIFS({crit})")
        ws.cell(r, 4, f'=COUNTIFS({crit},Cases!$D$2:$D${last},"PASS")')
        ws.cell(r, 5, f"=C{r}-D{r}")
        ws.cell(r, 6, f'=IF(C{r}=0,"",D{r}/C{r})')
        ws.cell(r, 7, f'=IFERROR(AVERAGEIFS(Cases!$E$2:$E${last},{crit}),"")')
        ws.cell(r, 6).number_format = "0.0%"
        ws.cell(r, 7).number_format = "0.000"
        for c in range(1, 8):
            ws.cell(r, c).font = st["body"]
    note_row = len(pairs) + 3
    ws.cell(note_row, 1, "Counts are computed from the Cases sheet. Pass/fail already accounts for direction (toxicity is lower-is-better), "
                         "and ERROR rows (judge failed) count as not passed.").font = st["muted"]
    for col, w in zip("ABCDEFG", [12, 34, 9, 9, 9, 10, 11]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def _sheet_ops(wb, n_metrics, targets):
    st = _style()
    ws = wb.create_sheet("Ops report", 3)
    last = max(n_metrics + 1, 2)

    def val(metric_id):        # pull one metric from the Metrics sheet; 'n/a' if the run didn't measure it
        crit = f'Metrics!$B$2:$B${last},"{metric_id}"'
        return f'=IF(COUNTIFS({crit})=0,"n/a",SUMIFS(Metrics!$E$2:$E${last},{crit}))'

    ws["A1"] = "Operational report"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)

    # --- latency table
    ws["A3"] = "Latency (milliseconds)"
    ws["A3"].font = st["bold"]
    _header(ws, 4, ["Stage", "Mean", "p50", "p95", "p99", "Min", "Max"])
    stages = [("End-to-end", "e2e"), ("Time to first token", "ttft"), ("Retrieval", "retrieval"), ("Generation", "generation")]
    for r, (label, key) in enumerate(stages, start=5):
        ws.cell(r, 1, label).font = st["bold"]
        for c, stat in enumerate(["mean", "p50", "p95", "p99", "min", "max"], start=2):
            cell = ws.cell(r, c, val(f"ops.latency.{key}_{stat}_ms"))
            cell.number_format, cell.font = "#,##0", st["body"]

    # --- targets (assumptions; edit the yellow cells to re-judge)
    ws["A10"] = "Targets (from evals/eval_ops.py -- edit the yellow cells to re-judge)"
    ws["A10"].font = st["bold"]
    _header(ws, 11, ["Check", "Target", "Measured", "Result"])
    checks = [
        ("End-to-end p95 (ms) <=", targets.get("slo_e2e_p95_ms"), "D5", "#,##0"),
        ("First-token p95 (ms) <=", targets.get("slo_ttft_p95_ms"), "D6", "#,##0"),
    ]
    r = 12
    for label, target, measured, fmt in checks:
        ws.cell(r, 1, label).font = st["body"]
        t = ws.cell(r, 2, target)
        t.fill, t.font, t.number_format = st["input_fill"], st["body"], fmt
        t.comment = Comment("Copied from evals/eval_ops.py; change it here to see how the run would score against another target.", "tracker")
        ws.cell(r, 3, f"={measured}").number_format = fmt
        ws.cell(r, 4, f'=IF(ISNUMBER(C{r}),IF(C{r}<=B{r},"PASS","FAIL"),"n/a")').font = st["bold"]
        r += 1

    # --- cost
    r += 1
    ws.cell(r, 1, "Cost").font = st["bold"]
    r += 1
    _header(ws, r, ["Metric", "Value"])
    cost_rows = [("Cost per query (USD)", "cost_per_query_usd", "$0.000000"), ("Cost per query (INR)", "cost_per_query_inr", "0.0000"),
                 ("Avg input tokens", "avg_input_tokens", "#,##0"), ("Avg output tokens", "avg_output_tokens", "#,##0"),
                 ("Avg cached tokens", "avg_cached_tokens", "#,##0"), ("Output share of cost (%)", "output_cost_share_pct", "0.0"),
                 ("Projected monthly (USD)", "monthly_usd", "$#,##0.00")]
    cost_first = r + 1
    for label, key, fmt in cost_rows:
        r += 1
        ws.cell(r, 1, label).font = st["body"]
        cell = ws.cell(r, 2, val(f"ops.cost.{key}"))
        cell.number_format, cell.font = fmt, st["body"]
    r += 1
    ws.cell(r, 1, "Budget per query (USD) <=").font = st["body"]
    b = ws.cell(r, 2, targets.get("cost_budget_usd"))
    b.fill, b.font, b.number_format = st["input_fill"], st["body"], "$0.000000"
    b.comment = Comment("Copied from evals/eval_ops.py (COST_BUDGET_PER_QUERY_USD).", "tracker")
    ws.cell(r, 3, f"=B{cost_first}").number_format = "$0.000000"
    ws.cell(r, 4, f'=IF(ISNUMBER(C{r}),IF(C{r}<=B{r},"PASS","FAIL"),"n/a")').font = st["bold"]
    budget_row = r

    # --- reliability
    r += 2
    ws.cell(r, 1, "Reliability").font = st["bold"]
    r += 1
    _header(ws, r, ["Metric", "Value"])
    for label, key, fmt in [("Total requests", "total_requests", "#,##0"), ("Success rate (%)", "success_rate", "0.00"),
                            ("Error rate (%)", "error_rate", "0.00"), ("Retry rate (%)", "retry_rate", "0.00")]:
        r += 1
        ws.cell(r, 1, label).font = st["body"]
        cell = ws.cell(r, 2, val(f"ops.reliability.{key}"))
        cell.number_format, cell.font = fmt, st["body"]

    _pass_fail_formatting(ws, f"D12:D{budget_row}")
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 13
